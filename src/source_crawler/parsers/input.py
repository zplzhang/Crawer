from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Iterator, Sequence

from ..models import CandidateRules, Seed
from ..validators.urls import DOWNLOAD_EXTENSIONS, extract_urls, normalize_url


def load_job_settings(input_path: str | None) -> dict[str, object]:
    """读取 JSON 顶层 run 配置；命令行参数仍可覆盖这些值。"""
    if not input_path or Path(input_path).suffix.lower() != ".json":
        return {}
    payload = json.loads(Path(input_path).read_text(encoding="utf-8-sig"))
    if not isinstance(payload, dict):
        return {}
    run = payload.get("run", {})
    if not isinstance(run, dict):
        raise ValueError("run 必须是对象")
    settings = dict(run)
    job = payload.get("job", {})
    if job is not None and not isinstance(job, dict):
        raise ValueError("job 必须是对象")
    if "max_candidates" not in settings and isinstance(job, dict) and "latest_items" in job:
        settings["max_candidates"] = job["latest_items"]
    return settings


def _string_list(value: object, field: str) -> tuple[str, ...]:
    if not isinstance(value, list) or any(not isinstance(item, str) or not item.strip() for item in value):
        raise ValueError(f"crawl.{field} 必须是非空字符串组成的数组")
    return tuple(item.strip() for item in value)


def _optional_int(value: object, field: str, minimum: int) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, int) or value < minimum:
        raise ValueError(f"crawl.{field} 必须是大于等于 {minimum} 的整数")
    return value


def parse_candidate_rules(value: object) -> CandidateRules:
    if value is None:
        value = {}
    if not isinstance(value, dict):
        raise ValueError("crawl 必须是对象")
    known = {"max_depth", "max_pages", "max_items", "min_link_score", "include_terms", "exclude_terms",
             "region_terms", "year_start", "year_end", "dataset_detail_url_patterns", "scope_ready",
             "allowed_domains", "file_types", "download", "discovery_urls"}
    unknown = set(value) - known
    if unknown:
        raise ValueError(f"crawl 含未知字段: {', '.join(sorted(unknown))}")
    allowed_domains = tuple(domain.lower().strip(".") for domain in _string_list(value.get("allowed_domains", []), "allowed_domains"))
    if any(not re.fullmatch(r"[a-z0-9.-]+", domain) or ".." in domain for domain in allowed_domains):
        raise ValueError("crawl.allowed_domains 只能填写域名，不包含协议或路径")
    file_types = tuple(item.lower().lstrip(".") for item in _string_list(value.get("file_types", []), "file_types"))
    if any("." + item not in DOWNLOAD_EXTENSIONS for item in file_types):
        raise ValueError("crawl.file_types 含不支持的文件类型")
    download = value.get("download", True)
    if not isinstance(download, bool):
        raise ValueError("crawl.download 必须是 true 或 false")
    scope_ready = value.get("scope_ready", False)
    if not isinstance(scope_ready, bool):
        raise ValueError("crawl.scope_ready 必须是 true 或 false")
    year_start = _optional_int(value.get("year_start"), "year_start", 1000)
    year_end = _optional_int(value.get("year_end"), "year_end", 1000)
    if year_start is not None and year_end is not None and year_start > year_end:
        raise ValueError("crawl.year_start 不能大于 year_end")
    discovery_urls = tuple(normalize_url(url) for url in _string_list(
        value.get("discovery_urls", []), "discovery_urls"
    ))
    if any(not url for url in discovery_urls):
        raise ValueError("crawl.discovery_urls 必须是有效网址")
    return CandidateRules(
        max_depth=_optional_int(value.get("max_depth"), "max_depth", 0),
        max_pages=_optional_int(value.get("max_pages"), "max_pages", 1),
        max_items=_optional_int(value.get("max_items"), "max_items", 1),
        min_link_score=_optional_int(value.get("min_link_score"), "min_link_score", 0),
        include_terms=_string_list(value.get("include_terms", []), "include_terms"),
        exclude_terms=_string_list(value.get("exclude_terms", []), "exclude_terms"),
        region_terms=_string_list(value.get("region_terms", []), "region_terms"),
        year_start=year_start,
        year_end=year_end,
        dataset_detail_url_patterns=_string_list(value.get("dataset_detail_url_patterns", []), "dataset_detail_url_patterns"),
        discovery_urls=discovery_urls,
        scope_ready=scope_ready,
        allowed_domains=allowed_domains,
        file_types=file_types,
        download=download,
    )

def iter_input_seeds(
    urls: Sequence[str],
    input_path: str | None,
    sheet_name: str | None,
    header_row: int,
    url_column: str | None,
) -> Iterator[Seed]:
    seen: set[str] = set()
    for raw in urls:
        for url in extract_urls(raw) or [normalize_url(raw)]:
            if url and url not in seen:
                seen.add(url)
                yield Seed(url=url)
    if not input_path:
        return
    path = Path(input_path)
    if not path.exists():
        raise FileNotFoundError(path)
    suffix = path.suffix.lower()
    if suffix in {".txt", ".list"}:
        for line in path.read_text(encoding="utf-8-sig").splitlines():
            for url in extract_urls(line) or [normalize_url(line)]:
                if url and url not in seen:
                    seen.add(url)
                    yield Seed(url=url)
        return
    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
        candidates = payload.get("candidates", []) if isinstance(payload, dict) else payload
        defaults: dict[str, object] = {}
        if isinstance(payload, dict):
            job = payload.get("job", {})
            if job is not None and not isinstance(job, dict):
                raise ValueError("job 必须是对象")
            if isinstance(job, dict):
                defaults = {
                    "name": job.get("material_or_data", ""),
                    "country_or_region": job.get("target_region", ""),
                    "first_batch_scope": job.get("first_batch_scope", ""),
                    "license_and_boundaries": job.get("license_and_boundaries", ""),
                }
            configured_defaults = payload.get("candidate_defaults", {})
            if not isinstance(configured_defaults, dict):
                raise ValueError("candidate_defaults 必须是对象")
            defaults.update(configured_defaults)
            if candidates == []:
                source_urls = extract_urls(job.get("source_url", "")) if isinstance(job, dict) else []
                if len(source_urls) != 1:
                    raise ValueError("candidates 为空时，job.source_url 必须配置恰好一个有效网址")
                candidates = [{"id": "SITE-ENTRY", "url": source_urls[0]}]
        if not isinstance(candidates, list):
            raise ValueError("JSON 入口应为列表或包含 candidates 列表")
        candidate_ids: set[str] = set()
        json_urls: set[str] = set()
        for index, candidate in enumerate(candidates, start=1):
            if isinstance(candidate, str):
                candidate = {"url": candidate}
            if not isinstance(candidate, dict):
                raise ValueError("每个 candidate 应为网址或对象")
            default_crawl = defaults.get("crawl", {})
            candidate_crawl = candidate.get("crawl", {})
            if not isinstance(default_crawl, dict) or not isinstance(candidate_crawl, dict):
                raise ValueError("candidate_defaults.crawl 和 candidate.crawl 必须是对象")
            candidate = {
                **defaults,
                **candidate,
                "crawl": {**default_crawl, **candidate_crawl},
            }
            candidate_id = str(candidate.get("id") or candidate.get("source_id") or f"candidate-{index}")
            if candidate_id in candidate_ids:
                raise ValueError(f"候选编号重复: {candidate_id}；请为每个网址设置唯一 id")
            candidate_ids.add(candidate_id)
            enabled = candidate.get("enabled", True)
            if not isinstance(enabled, bool):
                raise ValueError(f"候选 {candidate_id} 的 enabled 必须是 true 或 false")
            rules = parse_candidate_rules(candidate.get("crawl"))
            if not enabled:
                continue
            candidate_urls = extract_urls(candidate.get("url", ""))
            if len(candidate_urls) != 1:
                raise ValueError(f"候选 {candidate_id} 必须配置恰好一个有效网址")
            url = candidate_urls[0]
            if url in json_urls:
                raise ValueError(f"候选 {candidate_id} 的网址重复: {url}")
            json_urls.add(url)
            if url in seen:
                continue
            seen.add(url)
            yield Seed(
                url=url,
                candidate_id=candidate_id,
                source_id=str(candidate.get("source_id", "") or ""),
                priority=str(candidate.get("priority", "") or ""),
                name=str(candidate.get("name", "") or ""),
                country_or_region=str(candidate.get("country_or_region", "") or ""),
                publisher=str(candidate.get("publisher", "") or ""),
                publication_date=str(candidate.get("publication_date", "") or ""),
                usage_license=str(candidate.get("usage_license", "") or ""),
                first_batch_scope=str(candidate.get("first_batch_scope", "") or ""),
                license_and_boundaries=str(candidate.get("license_and_boundaries", "") or ""),
                rules=rules,
            )
        return
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if not reader.fieldnames:
                return
            url_key = url_column or next((key for key in reader.fieldnames if key.lower() in {"url", "网址", "链接", "来源入口", "原始网址"}), reader.fieldnames[0])
            for index, row in enumerate(reader, start=2):
                for url in extract_urls(row.get(url_key, "")):
                    if url not in seen:
                        seen.add(url)
                        yield Seed(
                            url=url,
                            candidate_id=str(row.get("id", row.get("编号", index)) or index),
                            source_id=str(row.get("编号", row.get("id", index)) or index),
                            priority=str(row.get("优先级", "") or ""),
                            name=str(row.get("材料或数据", row.get("名称", row.get("name", ""))) or ""),
                            country_or_region=str(row.get("区域", row.get("国家或地区", row.get("region", ""))) or ""),
                            first_batch_scope=str(row.get("首批范围", "") or ""),
                            license_and_boundaries=str(row.get("许可与边界", "") or ""),
                        )
        return
    if suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("读取 Excel 需要安装 openpyxl") from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
        headers = {str(sheet.cell(header_row, col).value or "").strip(): col for col in range(1, sheet.max_column + 1)}
        aliases = (url_column,) if url_column else ("来源入口", "原始网址", "网址", "链接", "URL", "url")
        url_col = next((headers[name] for name in aliases if name and name in headers), None)
        if not url_col:
            raise ValueError(f"未找到链接列。现有表头: {', '.join(k for k in headers if k)}")
        id_col = next((headers[name] for name in ("编号", "ID", "id") if name in headers), None)
        name_col = next((headers[name] for name in ("材料或数据", "名称", "name") if name in headers), None)
        region_col = next((headers[name] for name in ("区域", "国家或地区", "region") if name in headers), None)
        priority_col = headers.get("优先级")
        scope_col = headers.get("首批范围")
        boundary_col = headers.get("许可与边界")
        id_counts: dict[str, int] = {}
        for row in range(header_row + 1, sheet.max_row + 1):
            for url in extract_urls(sheet.cell(row, url_col).value):
                if url in seen:
                    continue
                seen.add(url)
                source_id = str(sheet.cell(row, id_col).value or row) if id_col else str(row)
                id_counts[source_id] = id_counts.get(source_id, 0) + 1
                yield Seed(
                    url=url,
                    candidate_id=source_id if id_counts[source_id] == 1 else f"{source_id}-{id_counts[source_id]}",
                    source_id=source_id,
                    priority=str(sheet.cell(row, priority_col).value or "") if priority_col else "",
                    name=str(sheet.cell(row, name_col).value or "") if name_col else "",
                    country_or_region=str(sheet.cell(row, region_col).value or "") if region_col else "",
                    first_batch_scope=str(sheet.cell(row, scope_col).value or "") if scope_col else "",
                    license_and_boundaries=str(sheet.cell(row, boundary_col).value or "") if boundary_col else "",
                )
        return
    raise ValueError(f"不支持的输入格式: {suffix}")
