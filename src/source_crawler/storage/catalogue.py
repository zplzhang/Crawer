from __future__ import annotations

import json
import sqlite3
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from typing import Iterable, Iterator

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import OUTPUT_COLUMNS, Record, utc_now


WORKBOOK_NAME = "采集结果.xlsx"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _excel_value(field: str, value: object) -> object:
    if not isinstance(value, str):
        return value
    if field == "publication_date" and len(value) == 10:
        try:
            return datetime.strptime(value, "%Y-%m-%d")
        except ValueError:
            pass
    if field == "download_date" and value:
        try:
            return datetime.fromisoformat(value).replace(tzinfo=None)
        except ValueError:
            pass
    return "'" + value if value.startswith("=") else value


def _write_sheet(workbook: Workbook, title: str, fields: list[str], rows: list[dict[str, object]]) -> None:
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(fields)
    for row in rows:
        sheet.append([_excel_value(field, row.get(field, "")) for field in fields])
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24
    sheet.auto_filter.ref = sheet.dimensions
    date_fields = {"publication_date", "download_date"}
    for index, field in enumerate(fields, start=1):
        letter = get_column_letter(index)
        values = [str(row.get(field, "") or "") for row in rows[:200]]
        width = min(max([len(field), *(len(value) for value in values)]), 45) + 2
        sheet.column_dimensions[letter].width = max(width, 12)
        if field in date_fields and sheet.max_row > 1:
            fmt = "yyyy-mm-dd" if field == "publication_date" else "yyyy-mm-dd hh:mm:ss"
            for cell in sheet[letter][1:]:
                cell.number_format = fmt
        if field in {"original_url", "seed_url"}:
            for cell in sheet[letter][1:]:
                if isinstance(cell.value, str) and cell.value.startswith(("http://", "https://")):
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

class StateStore:
    def __init__(self, path: Path):
        self.connection = sqlite3.connect(path)
        self.connection.execute(
            "CREATE TABLE IF NOT EXISTS url_state (url TEXT PRIMARY KEY, etag TEXT, last_modified TEXT, sha256 TEXT, local_path TEXT, checked_at TEXT)"
        )
        self.connection.execute(
            "CREATE TABLE IF NOT EXISTS records (url TEXT PRIMARY KEY, payload TEXT NOT NULL, updated_at TEXT NOT NULL)"
        )
        self.connection.execute(
            "CREATE TABLE IF NOT EXISTS candidate_records (candidate_id TEXT NOT NULL, url TEXT NOT NULL, "
            "payload TEXT NOT NULL, updated_at TEXT NOT NULL, PRIMARY KEY(candidate_id, url))"
        )
        for url, payload, updated_at in self.connection.execute("SELECT url, payload, updated_at FROM records"):
            candidate_id = str(json.loads(payload).get("candidate_id", "") or "")
            self.connection.execute(
                "INSERT OR IGNORE INTO candidate_records(candidate_id, url, payload, updated_at) VALUES(?,?,?,?)",
                (candidate_id, url, payload, updated_at),
            )
        self.connection.commit()

    def get_url_state(self, url: str) -> dict[str, str]:
        row = self.connection.execute(
            "SELECT etag, last_modified, sha256, local_path FROM url_state WHERE url = ?", (url,)
        ).fetchone()
        if not row:
            return {}
        return {"etag": row[0] or "", "last_modified": row[1] or "", "sha256": row[2] or "", "local_path": row[3] or ""}

    def save_url_state(self, url: str, headers: requests.structures.CaseInsensitiveDict, sha256: str = "", local_path: str = "") -> None:
        previous = self.get_url_state(url)
        self.connection.execute(
            "INSERT INTO url_state(url, etag, last_modified, sha256, local_path, checked_at) VALUES(?,?,?,?,?,?) "
            "ON CONFLICT(url) DO UPDATE SET etag=excluded.etag, last_modified=excluded.last_modified, "
            "sha256=excluded.sha256, local_path=excluded.local_path, checked_at=excluded.checked_at",
            (
                url,
                headers.get("ETag", previous.get("etag", "")),
                headers.get("Last-Modified", previous.get("last_modified", "")),
                sha256 or previous.get("sha256", ""),
                local_path or previous.get("local_path", ""),
                utc_now(),
            ),
        )
        self.connection.commit()

    def save_record(self, record: Record) -> None:
        payload = json.dumps(asdict(record), ensure_ascii=False)
        self.connection.execute(
            "INSERT INTO candidate_records(candidate_id, url, payload, updated_at) VALUES(?,?,?,?) "
            "ON CONFLICT(candidate_id, url) DO UPDATE SET payload=excluded.payload, updated_at=excluded.updated_at",
            (record.candidate_id, record.original_url, payload, utc_now()),
        )
        self.connection.commit()

    def get_record(self, url: str, candidate_id: str = "") -> Record | None:
        row = self.connection.execute(
            "SELECT payload FROM candidate_records WHERE candidate_id = ? AND url = ?",
            (candidate_id, url),
        ).fetchone()
        return Record(**json.loads(row[0])) if row else None

    def records(self) -> Iterator[dict[str, str]]:
        for (payload,) in self.connection.execute(
            "SELECT payload FROM candidate_records ORDER BY updated_at, candidate_id, url"
        ):
            yield json.loads(payload)

    def export_records(self, output_dir: Path) -> None:
        rows = list(self.records())
        workbook = Workbook()
        del workbook[workbook.sheetnames[0]]
        _write_sheet(workbook, "采集记录", OUTPUT_COLUMNS, rows)
        workbook.save(output_dir / WORKBOOK_NAME)
        with (output_dir / "records.jsonl").open("w", encoding="utf-8") as handle:
            for row in rows:
                handle.write(json.dumps(row, ensure_ascii=False) + "\n")

    def export_progress(self, output_dir: Path, rows: Iterable[dict[str, object]]) -> None:
        fields = ["candidate_id", "source_id", "seed_url", "target_items", "items_collected",
                  "pages_processed", "stop_reason"]
        path = output_dir / WORKBOOK_NAME
        workbook = load_workbook(path) if path.exists() else Workbook()
        if workbook.sheetnames == ["Sheet"] and workbook["Sheet"].max_row == 1:
            del workbook["Sheet"]
        _write_sheet(workbook, "候选进度", fields, list(rows))
        workbook.save(path)

    def close(self) -> None:
        self.connection.close()
