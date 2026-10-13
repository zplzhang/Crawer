import json
import contextlib
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import requests
from openpyxl import load_workbook

from source_crawler import cli, config, models, pipeline
from source_crawler.parsers import input as source_input, page as metadata
from source_crawler.parsers.api import discover_json_links
from source_crawler.parsers.sru import discover_sru_links
from source_crawler.validators import urls
from source_crawler.validators.scope import is_pagination_link, scope_mismatch


class CrawlerTests(unittest.TestCase):
    def test_normalize_url_removes_fragment_and_tracking(self):
        url = urls.normalize_url("HTTPS://Example.org//data/?utm_source=x&id=3#part")
        self.assertEqual(url, "https://example.org/data/?id=3")

    def test_fao_subdomains_are_excluded(self):
        self.assertTrue(urls.is_excluded("https://data.fao.org/catalog", ["fao.org"]))
        self.assertFalse(urls.is_excluded("https://example.org/fao.org", ["fao.org"]))
        self.assertFalse(urls.within_site_scope("https://other.example.org/file.pdf", "https://data.example.org", ()))
        self.assertTrue(urls.within_site_scope("https://other.example.org/file.pdf", "https://data.example.org", ("other.example.org",)))

    def test_detect_extension_uses_url_content_type_and_disposition(self):
        self.assertEqual(urls.detect_extension("https://example.org/a/report.pdf"), ".pdf")
        self.assertEqual(urls.detect_extension("https://example.org/download?id=1", "text/csv"), ".csv")
        self.assertEqual(
            urls.detect_extension(
                "https://example.org/download", "application/octet-stream", 'attachment; filename="data.xlsx"'
            ),
            ".xlsx",
        )

    def test_requested_document_formats_are_downloadable(self):
        for extension in ("pdf", "doc", "docx", "txt", "md"):
            with self.subTest(extension=extension):
                self.assertEqual(urls.detect_extension(f"https://example.org/files/report.{extension}"),
                                 f".{extension}")
                rules = source_input.parse_candidate_rules({"file_types": [extension.upper()]})
                self.assertEqual(rules.file_types, (extension,))
        self.assertEqual(urls.detect_extension("https://example.org/download?id=1", "text/plain"), ".txt")
        self.assertEqual(urls.detect_extension("https://example.org/download?id=2", "text/markdown"), ".md")

    def test_extract_metadata_separates_blog_and_extracts_requested_fields(self):
        html = """
        <html><head>
          <title>Agriculture policy commentary 2015-2020</title>
          <meta name="citation_publisher" content="Example Institute">
          <meta name="citation_publication_date" content="2024-05-03">
          <meta name="license" content="CC BY 4.0">
        </head><body><p>Blog opinion. Updated: 2025-01-01.</p></body></html>
        """
        record, _ = metadata.extract_metadata(html, "https://example.org/blog/policy")
        self.assertEqual(record.publisher, "Example Institute")
        self.assertEqual(record.publication_date, "2024-05-03")
        self.assertEqual(record.coverage_years, "2015-2020")
        self.assertEqual(record.usage_license, "CC BY 4.0")
        self.assertEqual(record.source_category, "研究博客/观点文章")

    def test_discover_links_prioritizes_downloads_and_ignores_mailto(self):
        html = """
        <a href="/files/data.csv">Download data</a>
        <a href="/about">About us</a>
        <a href="mailto:test@example.org">Contact</a>
        """
        soup = metadata.BeautifulSoup(html, "html.parser")
        links = metadata.discover_links(soup, "https://example.org/topic")
        self.assertEqual(links[0][0], "https://example.org/files/data.csv")
        self.assertTrue(all(not url.startswith("mailto:") for url, _, _ in links))

    def test_dutch_official_data_links_and_pagination_are_relevant(self):
        self.assertGreaterEqual(metadata.score_link(
            "https://www.rvo.nl/onderwerpen/mest/tabellen", "Tabellen stikstof en mest"
        ), 3)
        self.assertTrue(is_pagination_link("https://example.org/archief?page=2", "Volgende"))

    def test_json_search_results_keep_title_date_and_next_page(self):
        links = discover_json_links({
            "searchResults": [{"title": "Mestnormen 2026", "created": "2026-01-01",
                               "urlAlias": "/files/mestnormen-2026.pdf"}],
            "pager": {"nextPage": 1},
        }, "https://www.rvo.nl/api/rvo/v1/search-publications?search=mest")
        by_url = {link.url: link for link in links}
        file_link = by_url["https://www.rvo.nl/files/mestnormen-2026.pdf"]
        self.assertEqual(file_link.title, "Mestnormen 2026")
        self.assertEqual(file_link.publication_date, "2026-01-01")
        self.assertTrue(any("page=1" in url for url in by_url))

    def test_discovery_urls_are_loaded_with_candidate_rules(self):
        rules = source_input.parse_candidate_rules({
            "discovery_urls": ["https://example.org/api/search?q=agriculture"],
        })
        self.assertEqual(rules.discovery_urls, ("https://example.org/api/search?q=agriculture",))

    def test_sru_results_keep_metadata_and_generate_next_batch(self):
        xml = b'''<?xml version="1.0"?><sru:searchRetrieveResponse xmlns:sru="urn:sru"
            xmlns:dcterms="http://purl.org/dc/terms/"><sru:numberOfRecords>201</sru:numberOfRecords>
            <sru:records><sru:record><dcterms:title>Mestbeleid 2026</dcterms:title>
            <dcterms:available>2026-01-02</dcterms:available>
            <dcterms:hasVersion resourceIdentifier="https://example.org/mest-2026.html"/>
            </sru:record></sru:records></sru:searchRetrieveResponse>'''
        links = discover_sru_links(xml, "https://repository.example/sru?operation=searchRetrieve&maximumRecords=100&startRecord=1")
        by_url = {link.url: link for link in links}
        self.assertEqual(by_url["https://example.org/mest-2026.html"].publication_date, "2026-01-02")
        self.assertTrue(any("startRecord=101" in url for url in by_url))

    def test_text_input_is_deduplicated(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "urls.txt"
            source.write_text("https://example.org/data\nhttps://example.org/data#top\n", encoding="utf-8")
            seeds = list(source_input.iter_input_seeds([], str(source), None, 1, None))
        self.assertEqual([seed.url for seed in seeds], ["https://example.org/data"])

    def test_default_candidate_list_is_ready_for_one_command_start(self):
        self.assertTrue(config.DEFAULT_INPUT.is_file())
        self.assertEqual(config.DEFAULT_INPUT.name, "site_config.json")
        default_payload = json.loads(config.DEFAULT_INPUT.read_text(encoding="utf-8"))
        seeds = list(source_input.iter_input_seeds([], str(config.DEFAULT_INPUT), None, 1, None))
        self.assertEqual(len(seeds), 2)
        self.assertEqual(seeds[0].candidate_id, "NL-RVO-MEST")
        self.assertEqual(seeds[0].url, default_payload["job"]["source_url"])
        self.assertEqual(seeds[0].country_or_region, default_payload["job"]["target_region"])
        output = io.StringIO()
        with contextlib.redirect_stdout(output), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(cli.main(["--dry-run"]), 0)
        self.assertIn("NL-RVO-MEST", output.getvalue())
        self.assertTrue(seeds[0].first_batch_scope)
        self.assertEqual(seeds[0].license_and_boundaries, "")
        self.assertTrue(all(seed.rules.max_items == 10000 for seed in seeds))
        self.assertTrue(all(seed.rules.scope_ready for seed in seeds))
        output = io.StringIO()
        with contextlib.redirect_stdout(output), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(cli.main(["--candidate-id", "NL-RVO-MEST", "--dry-run"]), 0)
        self.assertEqual(len(output.getvalue().splitlines()), 1)
        self.assertIn("NL-RVO-MEST", output.getvalue())

    def test_large_collection_requires_confirmed_scope(self):
        with self.assertLogs(level="WARNING"):
            self.assertEqual(cli.main(["--input", "config/candidates.json",
                                       "--candidate-id", "D02"]), 2)

    def test_empty_candidates_uses_job_source_url_and_job_metadata(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "site.json"
            path.write_text(json.dumps({
                "job": {
                    "source_url": "https://example.org/catalog",
                    "target_region": "东非",
                    "material_or_data": "农业数据",
                    "first_batch_scope": "2020年以来",
                },
                "candidate_defaults": {
                    "source_id": "AGRI",
                    "crawl": {"scope_ready": True, "max_items": 10000},
                },
                "candidates": [],
            }), encoding="utf-8")
            seeds = list(source_input.iter_input_seeds([], str(path), None, 1, None))
            self.assertEqual(len(seeds), 1)
            self.assertEqual(seeds[0].candidate_id, "SITE-ENTRY")
            self.assertEqual(seeds[0].url, "https://example.org/catalog")
            self.assertEqual(seeds[0].country_or_region, "东非")
            self.assertEqual(seeds[0].first_batch_scope, "2020年以来")
            self.assertEqual(seeds[0].rules.max_items, 10000)

    def test_run_max_candidates_limits_entries_for_one_click_run(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "site.json"
            path.write_text(json.dumps({
                "run": {"max_candidates": 2},
                "candidate_defaults": {"crawl": {"scope_ready": True, "max_items": 10000}},
                "candidates": [
                    {"id": "one", "url": "https://example.org/one"},
                    {"id": "two", "url": "https://example.org/two"},
                    {"id": "three", "url": "https://example.org/three"},
                ],
            }), encoding="utf-8")
            output = io.StringIO()
            with contextlib.redirect_stdout(output), contextlib.redirect_stderr(io.StringIO()):
                self.assertEqual(cli.main(["--input", str(path), "--dry-run"]), 0)
            lines = output.getvalue().splitlines()
            self.assertEqual(len(lines), 2)
            self.assertIn("one", lines[0])
            self.assertIn("two", lines[1])

    def test_selected_candidate_gets_its_own_exploration_budget(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sites.json"
            path.write_text(json.dumps({"candidates": [{"id": "sample",
                "url": "https://example.org/start", "crawl": {"scope_ready": True,
                "max_items": 10000, "max_pages": 50000}}]}), encoding="utf-8")
            fake_crawler = Mock()
            with patch.object(cli, "Crawler", return_value=fake_crawler) as crawler_class:
                self.assertEqual(cli.main(["--input", str(path), "--candidate-id", "sample",
                                           "--output", directory]), 0)
            configuration = crawler_class.call_args.args[0]
            self.assertEqual(configuration.max_pages, 50000)
            self.assertEqual(configuration.max_pages_per_site, 50000)
            fake_crawler.crawl.assert_called_once()

    def test_scope_matches_keywords_region_and_year(self):
        rules = models.CandidateRules(include_terms=("wheat",), region_terms=("Pakistan",),
                                      year_start=2020, year_end=2024)
        self.assertEqual(scope_mismatch(rules, "Pakistan wheat statistics 2022"), "")
        self.assertIn("地区", scope_mismatch(rules, "India wheat statistics 2022"))
        self.assertIn("年份", scope_mismatch(rules, "Pakistan wheat statistics 2018"))

    def test_candidate_rules_are_loaded_and_validated(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "candidates.json"
            path.write_text(json.dumps({"candidates": [{
                "id": "sample", "url": "https://example.org/start", "country_or_region": "欧盟",
                "publisher": "Example Agency", "publication_date": "2026-10-01", "usage_license": "CC BY 4.0",
                "first_batch_scope": "仅政策 PDF", "license_and_boundaries": "逐份核验许可",
                "crawl": {"max_depth": 2, "max_pages": 3, "include_terms": ["policy"],
                          "allowed_domains": ["files.example.net"], "file_types": ["PDF"], "download": False},
            }]}), encoding="utf-8")
            seed = next(source_input.iter_input_seeds([], str(path), None, 1, None))
            self.assertEqual(seed.candidate_id, "sample")
            self.assertEqual(seed.first_batch_scope, "仅政策 PDF")
            self.assertEqual(seed.publisher, "Example Agency")
            self.assertEqual(seed.publication_date, "2026-10-01")
            self.assertEqual(seed.usage_license, "CC BY 4.0")
            self.assertEqual(seed.rules.allowed_domains, ("files.example.net",))
            self.assertEqual(seed.rules.file_types, ("pdf",))
            self.assertFalse(seed.rules.download)
            path.write_text(json.dumps({"candidates": [{"id": "bad", "url": "https://example.org",
                           "crawl": {"max_pages": -1}}]}), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "max_pages"):
                list(source_input.iter_input_seeds([], str(path), None, 1, None))

    def test_json_job_settings_and_candidate_defaults_are_inherited(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "site.json"
            path.write_text(json.dumps({
                "job": {
                    "target_region": "目标国",
                    "material_or_data": "国别简报",
                    "first_batch_scope": "最新版本",
                    "latest_items": 10,
                },
                "run": {"output_dir": "/tmp/results", "include_fao": True, "delay": 2.0},
                "candidate_defaults": {
                    "source_id": "SOURCE",
                    "publisher": "Official Agency",
                    "usage_license": "CC BY 4.0",
                    "crawl": {"scope_ready": True, "file_types": ["pdf"]},
                },
                "candidates": [{"id": "one", "url": "https://example.org/report.pdf"}],
            }), encoding="utf-8")
            settings = source_input.load_job_settings(str(path))
            self.assertEqual(settings["output_dir"], "/tmp/results")
            self.assertEqual(settings["max_candidates"], 10)
            seed = next(source_input.iter_input_seeds([], str(path), None, 1, None))
            self.assertEqual(seed.name, "国别简报")
            self.assertEqual(seed.country_or_region, "目标国")
            self.assertEqual(seed.first_batch_scope, "最新版本")
            self.assertEqual(seed.source_id, "SOURCE")
            self.assertEqual(seed.publisher, "Official Agency")
            self.assertEqual(seed.rules.file_types, ("pdf",))
            self.assertTrue(seed.rules.scope_ready)

    def test_private_network_urls_are_rejected_without_dns_lookup(self):
        with self.assertRaisesRegex(ValueError, "非公网"):
            urls.ensure_public_url("http://127.0.0.1/admin", resolve_dns=False)

    def test_crawl_downloads_file_before_following_page_and_keeps_metadata(self):
        with tempfile.TemporaryDirectory() as directory:
            engine = pipeline.Crawler(config.CrawlConfig(
                output_dir=Path(directory),
                excluded_domains=("fao.org",),
                max_depth=1,
                max_pages=3,
                max_pages_per_site=3,
                delay=0,
                timeout=1,
                max_file_mb=1,
                min_link_score=2,
                respect_robots=False,
                verify_dns=False,
                user_agent="test-agent",
            ))
            html = b'<title>Official datasets</title><meta name="license" content="CC BY 4.0"><a href="/about-data">Dataset overview</a><a href="/files/data.csv">Download CSV</a>'
            seen = []

            def fake_request(url, headers=None, stream=True, **kwargs):
                seen.append(url)
                response = Mock()
                response.url = url
                response.status_code = 200
                response.headers = requests.structures.CaseInsensitiveDict()
                response.encoding = "utf-8"
                if url.endswith(".csv"):
                    response.headers["Content-Type"] = "text/csv"
                    response.iter_content.return_value = iter([b"year,value\n2024,1\n"])
                else:
                    response.headers["Content-Type"] = "text/html"
                    response.content = html if url.endswith("/topic") else b"<title>Overview</title>"
                return response

            engine.http.get = fake_request
            try:
                engine.crawl([models.Seed(url="https://example.org/topic", source_id="D01")])
                self.assertEqual(seen[:2], ["https://example.org/topic", "https://example.org/files/data.csv"])
                rows = list(engine.state.records())
                self.assertEqual(len(rows), 3)
                file_record = next(row for row in rows if row["original_url"].endswith(".csv"))
                self.assertEqual(file_record["download_status"], "已下载")
                self.assertTrue((Path(directory) / file_record["local_path"]).exists())
                page_record = next(row for row in rows if row["original_url"].endswith("/topic"))
                self.assertEqual(page_record["usage_license"], "CC BY 4.0")

                response_304 = Mock()
                response_304.url = "https://example.org/topic"
                response_304.status_code = 304
                response_304.headers = requests.structures.CaseInsensitiveDict()
                engine.http.get = lambda *args, **kwargs: response_304
                unchanged, _ = engine.process(models.Task(
                    url="https://example.org/topic",
                    seed=models.Seed(url="https://example.org/topic", publisher="Updated Agency",
                                     publication_date="2026-10-01", usage_license="CC BY 4.0"),
                ))
                self.assertEqual(unchanged.download_status, "未变化")
                self.assertEqual(unchanged.usage_license, "CC BY 4.0")
                self.assertEqual(unchanged.publisher, "Updated Agency")
                self.assertEqual(unchanged.publication_date, "2026-10-01")
            finally:
                engine.close()

    def test_redirect_to_excluded_domain_is_not_requested(self):
        with tempfile.TemporaryDirectory() as directory:
            engine = pipeline.Crawler(config.CrawlConfig(
                output_dir=Path(directory),
                excluded_domains=("fao.org",),
                max_depth=0,
                max_pages=1,
                max_pages_per_site=1,
                delay=0,
                timeout=1,
                max_file_mb=1,
                min_link_score=2,
                respect_robots=False,
                verify_dns=False,
                user_agent="test-agent",
            ))
            redirect = Mock()
            redirect.status_code = 302
            redirect.headers = {"Location": "https://www.fao.org/data"}
            engine.http.session.get = Mock(return_value=redirect)
            try:
                with self.assertRaisesRegex(ValueError, "排除域名"):
                    engine.http.get("https://example.org/redirect")
                self.assertEqual(engine.http.session.get.call_count, 1)
            finally:
                engine.close()

    def test_transient_file_recheck_failure_keeps_existing_download(self):
        with tempfile.TemporaryDirectory() as directory:
            engine = pipeline.Crawler(config.CrawlConfig(
                output_dir=Path(directory), excluded_domains=(), max_depth=0, max_pages=1,
                max_pages_per_site=1, delay=0, timeout=1, max_file_mb=1,
                min_link_score=0, respect_robots=False, verify_dns=False, user_agent="test-agent",
            ))
            url = "https://example.org/report.pdf"
            existing = Path(directory) / "downloads" / "report.pdf"
            existing.parent.mkdir(parents=True, exist_ok=True)
            existing.write_bytes(b"%PDF-existing")
            engine.state.save_url_state(
                url, requests.structures.CaseInsensitiveDict(), "abc123",
                str(existing.relative_to(directory)),
            )
            engine.http.get = Mock(side_effect=requests.ConnectionError("temporary disconnect"))
            try:
                record, _ = engine.process(models.Task(
                    url=url,
                    seed=models.Seed(url=url, candidate_id="pdf", publisher="Official Agency",
                                     publication_date="2026-10-01", usage_license="CC BY 4.0"),
                ))
                self.assertEqual(record.download_status, "未变化")
                self.assertEqual(record.item_type, "文件")
                self.assertEqual(record.scope_status, "范围内")
                self.assertEqual(record.local_path, "downloads/report.pdf")
                self.assertIn("已保留先前下载文件", record.failure_reason)
            finally:
                engine.close()

    def test_redirect_outside_candidate_domains_is_not_requested(self):
        with tempfile.TemporaryDirectory() as directory:
            engine = pipeline.Crawler(config.CrawlConfig(output_dir=Path(directory), delay=0,
                respect_robots=False, verify_dns=False))
            redirect = Mock()
            redirect.status_code = 302
            redirect.headers = {"Location": "https://other.example.net/file.pdf"}
            engine.http.session.get = Mock(return_value=redirect)
            try:
                with self.assertRaisesRegex(ValueError, "允许域名"):
                    engine.http.get("https://example.org/start", site_seed_url="https://example.org/start")
                self.assertEqual(engine.http.session.get.call_count, 1)
            finally:
                engine.close()

    def test_per_candidate_scope_limits_links_and_pages(self):
        with tempfile.TemporaryDirectory() as directory:
            engine = pipeline.Crawler(config.CrawlConfig(
                output_dir=Path(directory), max_pages=10, max_pages_per_site=10,
                delay=0, respect_robots=False, verify_dns=False,
            ))
            seen = []
            html = (b'<title>Official policy</title><a href="/files/keep-one.pdf">Policy keep</a>'
                    b'<a href="/files/keep-two.pdf">Policy keep</a>'
                    b'<a href="/files/skip.csv">Dataset</a>'
                    b'<a href="https://other.example.net/files/keep.pdf">Policy keep</a>')

            def fake_request(url, **kwargs):
                seen.append(url)
                response = Mock()
                response.url = url
                response.status_code = 200
                response.headers = requests.structures.CaseInsensitiveDict(
                    {"Content-Type": "application/pdf" if url.endswith(".pdf") else "text/html"}
                )
                response.encoding = "utf-8"
                response.content = html
                response.iter_content.return_value = iter([b"%PDF-1.4 test"])
                return response

            engine.http.get = fake_request
            seed = models.Seed(
                url="https://example.org/start", candidate_id="policy", source_id="D02",
                country_or_region="欧盟", first_batch_scope="政策 PDF", license_and_boundaries="核验许可",
                rules=models.CandidateRules(max_depth=1, max_pages=2, include_terms=("keep",),
                                             file_types=("pdf",)),
            )
            try:
                engine.crawl([seed])
                self.assertEqual(len(seen), 2)
                self.assertNotIn("https://other.example.net/files/keep.pdf", seen)
                self.assertTrue(seen[1].endswith(".pdf"))
                records = list(engine.state.records())
                self.assertEqual(len(records), 2)
                self.assertTrue(all(row["candidate_id"] == "policy" for row in records))
                self.assertTrue(all(row["first_batch_scope"] == "政策 PDF" for row in records))
            finally:
                engine.close()

    def test_record_only_configuration_does_not_request_file(self):
        with tempfile.TemporaryDirectory() as directory:
            engine = pipeline.Crawler(config.CrawlConfig(output_dir=Path(directory),
                delay=0, respect_robots=False, verify_dns=False))
            engine.http.get = Mock()
            seed = models.Seed(url="https://example.org/report.pdf", candidate_id="readonly",
                               rules=models.CandidateRules(download=False))
            try:
                engine.crawl([seed])
                engine.http.get.assert_not_called()
                record = next(engine.state.records())
                self.assertEqual(record["download_status"], "仅记录链接")
                self.assertFalse((Path(directory) / "downloads" / "report.pdf").exists())
            finally:
                engine.close()

    def test_file_type_configuration_skips_unlisted_file(self):
        with tempfile.TemporaryDirectory() as directory:
            engine = pipeline.Crawler(config.CrawlConfig(output_dir=Path(directory),
                delay=0, respect_robots=False, verify_dns=False))
            engine.http.get = Mock()
            seed = models.Seed(url="https://example.org/data.csv", candidate_id="pdf-only",
                               rules=models.CandidateRules(file_types=("pdf",)))
            try:
                record, _ = engine.process(models.Task(url=seed.url, seed=seed))
                engine.http.get.assert_not_called()
                self.assertEqual(record.download_status, "已跳过")
                self.assertIn("文件类型", record.failure_reason)
            finally:
                engine.close()

    def test_out_of_scope_file_is_not_downloaded_or_counted(self):
        with tempfile.TemporaryDirectory() as directory:
            engine = pipeline.Crawler(config.CrawlConfig(output_dir=Path(directory),
                delay=0, respect_robots=False, verify_dns=False))
            engine.http.get = Mock()
            seed = models.Seed(url="https://example.org/india-maize-2018.pdf", candidate_id="wheat",
                rules=models.CandidateRules(include_terms=("wheat",), region_terms=("Pakistan",),
                    year_start=2020, year_end=2024, max_items=10000, scope_ready=True))
            try:
                engine.crawl([seed])
                engine.http.get.assert_not_called()
                record = next(engine.state.records())
                self.assertEqual(record["scope_status"], "范围外")
                workbook = load_workbook(Path(directory) / "采集结果.xlsx", data_only=True)
                progress = list(workbook["候选进度"].values)
                self.assertEqual(progress[1][4], 0)
                self.assertFalse((Path(directory) / "records.csv").exists())
                self.assertFalse((Path(directory) / "candidate_progress.csv").exists())
            finally:
                engine.close()

    def test_shared_file_is_handled_by_each_candidate_rules(self):
        with tempfile.TemporaryDirectory() as directory:
            engine = pipeline.Crawler(config.CrawlConfig(output_dir=Path(directory),
                max_pages=10, max_pages_per_site=10, delay=0, respect_robots=False, verify_dns=False))
            requested = []

            def fake_request(url, **kwargs):
                requested.append(url)
                response = Mock()
                response.url = url
                response.status_code = 200
                response.headers = requests.structures.CaseInsensitiveDict(
                    {"Content-Type": "application/pdf" if url.endswith(".pdf") else "text/html"}
                )
                response.encoding = "utf-8"
                response.content = b'<title>Index</title><a href="/shared.pdf">Download report</a>'
                response.iter_content.return_value = iter([b"%PDF-1.4 test"])
                return response

            engine.http.get = fake_request
            seeds = [
                models.Seed(url="https://example.org/a", candidate_id="a",
                            rules=models.CandidateRules(max_depth=1, download=False)),
                models.Seed(url="https://example.org/b", candidate_id="b",
                            rules=models.CandidateRules(max_depth=1, download=True)),
            ]
            try:
                engine.crawl(seeds)
                self.assertEqual(requested.count("https://example.org/shared.pdf"), 1)
                file_rows = [row for row in engine.state.records() if row["original_url"].endswith("shared.pdf")]
                self.assertEqual(len(file_rows), 2)
                self.assertEqual({row["candidate_id"]: row["download_status"] for row in file_rows},
                                 {"a": "仅记录链接", "b": "已下载"})
            finally:
                engine.close()

    def test_max_items_counts_files_and_dataset_details_but_not_navigation(self):
        with tempfile.TemporaryDirectory() as directory:
            engine = pipeline.Crawler(config.CrawlConfig(output_dir=Path(directory),
                max_pages=10, max_pages_per_site=10, delay=0, respect_robots=False, verify_dns=False))
            requested = []

            def fake_request(url, **kwargs):
                requested.append(url)
                response = Mock()
                response.url = url
                response.status_code = 200
                response.headers = requests.structures.CaseInsensitiveDict(
                    {"Content-Type": "application/pdf" if url.endswith(".pdf") else "text/html"}
                )
                response.encoding = "utf-8"
                if url.endswith("/start"):
                    response.content = (b'<title>Index</title>'
                        b'<a href="/files/pakistan-wheat-2022.pdf">Pakistan wheat 2022 PDF</a>'
                        b'<a href="/dataset/pakistan-wheat-2022">Pakistan wheat 2022 dataset</a>'
                        b'<a href="/about-data">About data</a>')
                else:
                    response.content = b'<title>Pakistan wheat 2022 dataset</title>'
                response.iter_content.return_value = iter([b"%PDF-1.4 test"])
                return response

            engine.http.get = fake_request
            seed = models.Seed(url="https://example.org/start", candidate_id="wheat",
                rules=models.CandidateRules(max_depth=2, max_pages=10, max_items=2,
                    include_terms=("wheat",), region_terms=("Pakistan",),
                    year_start=2022, year_end=2022,
                    dataset_detail_url_patterns=("/dataset/",), scope_ready=True))
            try:
                engine.crawl([seed])
                self.assertEqual(len(requested), 3)
                rows = list(engine.state.records())
                self.assertEqual(len(rows), 3)
                self.assertEqual({row["item_type"] for row in rows if row["item_type"]},
                                 {"文件", "数据集详情页"})
                workbook = load_workbook(Path(directory) / "采集结果.xlsx", data_only=True)
                headers = [cell.value for cell in workbook["候选进度"][1]]
                values = [cell.value for cell in workbook["候选进度"][2]]
                progress = dict(zip(headers, values))
                self.assertEqual(progress["items_collected"], 2)
                self.assertEqual(progress["pages_processed"], 3)
                self.assertEqual(progress["stop_reason"], "达到条目目标")
            finally:
                engine.close()

    def test_failed_download_does_not_consume_item_target(self):
        with tempfile.TemporaryDirectory() as directory:
            engine = pipeline.Crawler(config.CrawlConfig(output_dir=Path(directory),
                max_pages=10, max_pages_per_site=10, delay=0, respect_robots=False, verify_dns=False))
            requested = []

            def fake_request(url, **kwargs):
                requested.append(url)
                response = Mock()
                response.url = url
                response.status_code = 404 if url.endswith("a-missing.pdf") else 200
                response.headers = requests.structures.CaseInsensitiveDict(
                    {"Content-Type": "application/pdf" if url.endswith(".pdf") else "text/html"}
                )
                response.encoding = "utf-8"
                response.content = (b'<a href="/a-missing.pdf">Missing</a>'
                                    b'<a href="/b-good.pdf">Good 1</a>'
                                    b'<a href="/c-good.pdf">Good 2</a>')
                response.iter_content.return_value = iter([b"%PDF-1.4 test"])
                return response

            engine.http.get = fake_request
            seed = models.Seed(url="https://example.org/start", candidate_id="target",
                rules=models.CandidateRules(max_depth=1, max_pages=10, max_items=2,
                    scope_ready=True))
            try:
                engine.crawl([seed])
                self.assertIn("https://example.org/a-missing.pdf", requested)
                self.assertIn("https://example.org/b-good.pdf", requested)
                self.assertIn("https://example.org/c-good.pdf", requested)
                workbook = load_workbook(Path(directory) / "采集结果.xlsx", data_only=True)
                headers = [cell.value for cell in workbook["候选进度"][1]]
                values = [cell.value for cell in workbook["候选进度"][2]]
                progress = dict(zip(headers, values))
                self.assertEqual(progress["items_collected"], 2)
                self.assertEqual(progress["stop_reason"], "达到条目目标")
            finally:
                engine.close()

    def test_pagination_does_not_consume_depth(self):
        with tempfile.TemporaryDirectory() as directory:
            engine = pipeline.Crawler(config.CrawlConfig(output_dir=Path(directory),
                max_pages=3, max_pages_per_site=3, delay=0, respect_robots=False, verify_dns=False))
            requested = []

            def fake_request(url, **kwargs):
                requested.append(url)
                response = Mock()
                response.url = url
                response.status_code = 200
                response.headers = requests.structures.CaseInsensitiveDict(
                    {"Content-Type": "application/pdf" if url.endswith(".pdf") else "text/html"}
                )
                response.encoding = "utf-8"
                response.content = (b'<a href="?page=2">Next</a>' if url.endswith("/start") else
                                    b'<a href="/report.pdf">Download report</a>')
                response.iter_content.return_value = iter([b"%PDF-1.4 test"])
                return response

            engine.http.get = fake_request
            try:
                engine.crawl([models.Seed(url="https://example.org/start", candidate_id="pages",
                    rules=models.CandidateRules(max_depth=1, max_pages=3, max_items=1))])
                self.assertEqual(requested, ["https://example.org/start",
                    "https://example.org/start?page=2", "https://example.org/report.pdf"])
            finally:
                engine.close()


if __name__ == "__main__":
    unittest.main()
